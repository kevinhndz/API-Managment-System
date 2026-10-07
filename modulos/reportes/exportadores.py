from io import BytesIO
from html import escape

from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet


COLUMNAS = {
    "matriculas": [("Cuenta", "cuenta"), ("Estudiante", "estudiante"), ("Carrera", "carrera"), ("Asignatura", "asignatura"), ("Seccion", "seccion"), ("Periodo", "periodo"), ("Estado", "estado"), ("Fecha", "fecha_matricula"), ("Nota final", "nota_final")],
    "secciones": [("Seccion", "seccion"), ("Asignatura", "asignatura"), ("Docente", "docente"), ("Periodo", "periodo"), ("Aula", "aula"), ("Cupo", "cupo_maximo"), ("Matriculados", "matriculados"), ("Estado", "estado")],
    "estudiantes": [("Cuenta", "cuenta"), ("Nombre", "nombre"), ("Correo", "correo"), ("Carrera", "carrera"), ("Activo", "estado")],
    "docentes": [("Empleado", "numero_empleado"), ("Nombres", "nombres"), ("Apellidos", "apellidos"), ("Correo", "correo"), ("Activo", "estado")],
    "calificaciones": [("Cuenta", "cuenta"), ("Estudiante", "estudiante"), ("Asignatura", "asignatura"), ("Seccion", "seccion"), ("Periodo", "periodo"), ("Primer parcial", "primer_parcial"), ("Segundo parcial", "segundo_parcial"), ("Tercer parcial", "tercer_parcial"), ("Nota final", "nota_final")],
}


def valor(item, campo):
    dato = getattr(item, campo, None)
    if dato is None:
        return ""
    if isinstance(dato, bool):
        return "Si" if dato else "No"
    return dato


def exportar_excel(modulo, registros):
    libro = Workbook()
    hoja = libro.active
    hoja.title = modulo.capitalize()
    columnas = COLUMNAS[modulo]
    hoja.append([titulo for titulo, _ in columnas])

    for item in registros:
        hoja.append([valor(item, campo) for _, campo in columnas])

    for celda in hoja[1]:
        celda.font = Font(color="FFFFFF", bold=True)
        celda.fill = PatternFill("solid", fgColor="5B0309")

    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions
    for columna in hoja.columns:
        ancho = max(len(str(celda.value or "")) for celda in columna)
        hoja.column_dimensions[columna[0].column_letter].width = min(ancho + 3, 50)

    archivo = BytesIO()
    libro.save(archivo)
    archivo.seek(0)
    return StreamingResponse(archivo, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="reporte_{modulo}.xlsx"'})


def exportar_pdf(modulo, registros):
    archivo = BytesIO()
    documento = SimpleDocTemplate(archivo, pagesize=landscape(A4), leftMargin=28, rightMargin=28)
    estilos = getSampleStyleSheet()
    columnas = COLUMNAS[modulo]
    ancho = (landscape(A4)[0] - 56) / len(columnas)
    filas = [[Paragraph(titulo, estilos["BodyText"]) for titulo, _ in columnas]]

    for item in registros:
        filas.append([Paragraph(escape(str(valor(item, campo))), estilos["BodyText"]) for _, campo in columnas])

    tabla = Table(filas, colWidths=[ancho] * len(columnas), repeatRows=1)
    tabla.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F4E8E9")),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D8C6C6")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
    ]))
    documento.build([Paragraph(f"Reporte de {modulo}", estilos["Title"]), Spacer(1, 14), tabla])
    archivo.seek(0)
    return StreamingResponse(archivo, media_type="application/pdf", headers={"Content-Disposition": f'attachment; filename="reporte_{modulo}.pdf"'})
