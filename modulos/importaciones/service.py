"""Importacion pequena con validacion previa y aislamiento por institucion."""

import csv
import hashlib
import io
import secrets
from datetime import date, datetime, timezone

from email_validator import EmailNotValidError, validate_email
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from modulos.importaciones import repository
from modulos.importaciones.schema import ErrorImportacion

COLUMNAS = {
    "carreras": ("codigo", "nombre", "duracion_anios"),
    "docentes": ("numero_empleado", "nombres", "apellidos", "correo", "especialidad"),
    "estudiantes": ("cuenta", "nombre", "correo", "carrera_codigo"),
}
MAX_BYTES = 2 * 1024 * 1024
MAX_FILAS = 5000


def _texto(valor) -> str:
    return str(valor).strip() if valor is not None else ""


def _filas(archivo: bytes, nombre: str) -> tuple[list[dict], list[str]]:
    if not archivo or len(archivo) > MAX_BYTES:
        raise ErrorImportacion(413, "El archivo debe tener entre 1 byte y 2 MB")
    if nombre.lower().endswith(".csv"):
        try:
            lector = csv.DictReader(io.StringIO(archivo.decode("utf-8-sig")))
            cabeceras = [_texto(col).casefold() for col in (lector.fieldnames or [])]
            if len(cabeceras) != len(set(cabeceras)) or not all(cabeceras):
                raise ErrorImportacion(422, "El CSV tiene columnas vacias o repetidas")
            filas = []
            for fila in lector:
                if len(filas) >= MAX_FILAS:
                    raise ErrorImportacion(413, "Para mas de 5000 filas solicita una migracion asistida")
                filas.append({_texto(clave).casefold(): valor for clave, valor in fila.items() if clave is not None})
            return filas, cabeceras
        except UnicodeDecodeError as exc:
            raise ErrorImportacion(422, "El CSV debe usar UTF-8") from exc
    if nombre.lower().endswith(".xlsx"):
        try:
            libro = load_workbook(io.BytesIO(archivo), read_only=True, data_only=True)
            hoja = libro.active
            iterador = hoja.iter_rows(values_only=True)
            cabeceras = [_texto(valor).casefold() for valor in next(iterador, ())]
            if len(cabeceras) != len(set(cabeceras)) or not all(cabeceras):
                raise ErrorImportacion(422, "El Excel tiene columnas vacias o repetidas")
            filas = []
            for fila in iterador:
                if not any(valor is not None for valor in fila):
                    continue
                if len(filas) >= MAX_FILAS:
                    raise ErrorImportacion(413, "Para mas de 5000 filas solicita una migracion asistida")
                filas.append(dict(zip(cabeceras, fila)))
            libro.close()
            return filas, cabeceras
        except HTTPException:
            raise
        except Exception as exc:
            raise ErrorImportacion(422, "No se pudo leer el Excel") from exc
    raise ErrorImportacion(422, "Usa un archivo CSV o XLSX")


def _validar(tipo: str, fila: dict, carreras: dict[str, int]) -> dict:
    dato = {clave: _texto(valor) for clave, valor in fila.items() if clave}
    for columna in COLUMNAS[tipo]:
        if not dato.get(columna):
            raise ValueError(f"Falta {columna}")
    if tipo == "carreras":
        if len(dato["codigo"]) > 20 or len(dato["nombre"]) > 150:
            raise ValueError("Codigo o nombre demasiado largo")
        try:
            duracion = int(dato["duracion_anios"])
        except ValueError as exc:
            raise ValueError("Duracion invalida") from exc
        if not 1 <= duracion <= 15:
            raise ValueError("Duracion fuera de rango")
        return {"codigo": dato["codigo"], "nombre": dato["nombre"], "duracion_anios": duracion}
    try:
        correo = validate_email(dato["correo"], check_deliverability=False).normalized.casefold()
    except EmailNotValidError as exc:
        raise ValueError("Correo invalido") from exc
    if len(correo) > 150:
        raise ValueError("Correo demasiado largo")
    if tipo == "docentes":
        limites = {"numero_empleado": 20, "nombres": 100, "apellidos": 100, "especialidad": 120,
                   "cargo": 120, "departamento": 120, "tipo_contrato": 80}
        if any(len(dato.get(campo, "")) > limite for campo, limite in limites.items()):
            raise ValueError("Uno de los campos excede su longitud")
        fecha = dato.get("fecha_ingreso")
        if fecha:
            try:
                fecha = date.fromisoformat(fecha[:10]).isoformat()
            except ValueError as exc:
                raise ValueError("fecha_ingreso debe usar AAAA-MM-DD") from exc
        return {**{campo: dato.get(campo) or None for campo in limites}, "correo": correo,
                "fecha_ingreso": fecha or None}
    if len(dato["cuenta"]) > 20 or len(dato["cuenta"]) < 5 or len(dato["nombre"]) > 150:
        raise ValueError("Cuenta o nombre invalido")
    carrera_id = carreras.get(dato["carrera_codigo"].casefold())
    if not carrera_id:
        raise ValueError("La carrera no existe en esta universidad")
    fecha = dato.get("fecha_nacimiento")
    if fecha:
        try:
            fecha = date.fromisoformat(fecha[:10]).isoformat()
        except ValueError as exc:
            raise ValueError("fecha_nacimiento debe usar AAAA-MM-DD") from exc
    telefono = dato.get("telefono") or None
    if telefono and len(telefono) > 20:
        raise ValueError("Telefono demasiado largo")
    return {"cuenta": dato["cuenta"], "nombre": dato["nombre"], "correo": correo,
            "carrera_id": carrera_id, "telefono": telefono, "fechaNacimiento": fecha or None}


def previsualizar(db: Session, tipo: str, archivo: bytes, nombre: str) -> dict:
    if tipo not in COLUMNAS:
        raise ErrorImportacion(422, "Tipo de importacion no admitido")
    filas, cabeceras = _filas(archivo, nombre)
    faltantes = set(COLUMNAS[tipo]) - set(cabeceras)
    if faltantes:
        raise ErrorImportacion(422, "Faltan columnas: " + ", ".join(sorted(faltantes)))
    if len(filas) > MAX_FILAS:
        raise ErrorImportacion(413, "Para mas de 5000 filas solicita una migracion asistida")
    existentes_db = repository.datos_existentes(db, tipo)
    carreras = existentes_db["carreras"]
    principal = repository.PRINCIPALES[tipo]
    existentes = existentes_db["principales"]
    nombres_carrera = existentes_db["nombres_carrera"]
    correos = existentes_db["correos"]
    telefonos = existentes_db["telefonos"]
    vistos, correos_vistos, nombres_vistos, telefonos_vistos = set(), set(), set(), set()
    validos, errores = [], []
    for numero, fila in enumerate(filas, 2):
        try:
            registro = _validar(tipo, fila, carreras)
            clave = _texto(registro[principal]).casefold()
            if clave in vistos or clave in existentes:
                raise ValueError(f"{principal} duplicado en esta universidad")
            if tipo == "carreras" and (registro["nombre"].casefold() in nombres_carrera or registro["nombre"].casefold() in nombres_vistos):
                raise ValueError("Nombre de carrera duplicado en esta universidad")
            if tipo != "carreras" and (registro["correo"] in correos or registro["correo"] in correos_vistos):
                raise ValueError("Correo duplicado en esta universidad")
            if tipo == "estudiantes" and registro["telefono"] and (registro["telefono"] in telefonos or registro["telefono"] in telefonos_vistos):
                raise ValueError("Telefono duplicado en esta universidad")
            vistos.add(clave)
            if tipo == "carreras":
                nombres_vistos.add(registro["nombre"].casefold())
            if tipo != "carreras":
                correos_vistos.add(registro["correo"])
            if tipo == "estudiantes" and registro["telefono"]:
                telefonos_vistos.add(registro["telefono"])
            validos.append(registro)
        except ValueError as exc:
            errores.append({"fila": numero, "error": str(exc)})
    token = secrets.token_urlsafe(32)
    repository.guardar_vista(db, hashlib.sha256(token.encode()).hexdigest(), tipo, validos, errores)
    return {"token": token, "tipo": tipo, "total": len(filas), "validos": len(validos),
            "invalidos": len(errores), "errores": errores[:100]}


def aplicar(db: Session, token: str) -> dict:
    registro = repository.buscar_vista(db, hashlib.sha256(token.encode()).hexdigest())
    if not registro:
        raise ErrorImportacion(404, "Vista previa no encontrada en esta universidad")
    vencimiento = registro.vence_en if registro.vence_en.tzinfo else registro.vence_en.replace(tzinfo=timezone.utc)
    if registro.aplicado_en or datetime.now(timezone.utc) >= vencimiento:
        raise ErrorImportacion(409, "Esta vista previa ya se uso o vencio")
    if not registro.registros:
        raise ErrorImportacion(409, "No hay filas validas para importar")
    if not repository.aplicar_vista(db, registro):
        raise ErrorImportacion(409, "Los datos cambiaron desde la vista previa. Genera otra vista previa")
    return {"importados": len(registro.registros), "omitidos": len(registro.errores)}
